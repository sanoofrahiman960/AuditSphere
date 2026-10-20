"""Render a report dict (see builders.py) to PDF or XLSX bytes."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY = colors.HexColor("#0f1b2d")
ACCENT = colors.HexColor("#2563eb")
SERIES_COLORS = [colors.HexColor("#2563eb"), colors.HexColor("#f59e0b"), colors.HexColor("#10b981"),
                 colors.HexColor("#ef4444")]
GRID = colors.HexColor("#d7dde6")
ZEBRA = colors.HexColor("#f5f7fa")


def _fmt(value):
    if value is None or value == "":
        return "—"
    if isinstance(value, float):
        return f"{value:.1f}"
    return str(value)


def _generated(report):
    try:
        return datetime.fromisoformat(report["generated_at"]).strftime("%d %b %Y %H:%M UTC")
    except (KeyError, ValueError):
        return report.get("generated_at", "")


# --------------------------------------------------------------------------- PDF

def _chart(chart, width):
    labels = chart["labels"]
    if not labels or not chart["series"]:
        return None
    height = 70 * mm
    drawing = Drawing(width, height)
    bc = VerticalBarChart()
    bc.x, bc.y = 30, 45
    bc.width, bc.height = width - 180, height - 70
    bc.data = [[v or 0 for v in s["values"]] for s in chart["series"]]
    bc.categoryAxis.categoryNames = [label[:18] for label in labels]
    bc.categoryAxis.labels.angle = 30 if len(labels) > 5 else 0
    bc.categoryAxis.labels.boxAnchor = "ne" if len(labels) > 5 else "n"
    bc.categoryAxis.labels.fontSize = 7
    bc.valueAxis.labels.fontSize = 7
    bc.valueAxis.valueMin = 0
    max_value = max((max(row) if row else 0) for row in bc.data) or 1
    bc.valueAxis.valueMax = max_value * 1.15
    bc.barSpacing = 1
    bc.groupSpacing = 6
    for i in range(len(bc.data)):
        bc.bars[i].fillColor = SERIES_COLORS[i % len(SERIES_COLORS)]
        bc.bars[i].strokeColor = None
    drawing.add(bc)
    drawing.add(String(30, height - 12, chart["title"], fontName="Helvetica-Bold", fontSize=9, fillColor=NAVY))
    legend = Legend()
    legend.x, legend.y = width - 140, height - 30
    legend.fontSize = 7
    legend.colorNamePairs = [(SERIES_COLORS[i % len(SERIES_COLORS)], s["name"]) for i, s in enumerate(chart["series"])]
    drawing.add(legend)
    return drawing


def render_pdf(report):
    buffer = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=16 * mm,
                            bottomMargin=14 * mm, title=report["title"], author="AuditSphere")
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=NAVY, fontSize=17, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=NAVY, fontSize=12, spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#475569"))
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7, leading=9, alignment=TA_LEFT)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")

    story = [
        Paragraph(report["organization"], small),
        Paragraph(report["title"], h1),
        Paragraph(f"Generated {_generated(report)} by {report.get('generated_by', '')}", small),
        Spacer(1, 4 * mm),
    ]

    filter_rows = [[Paragraph(f"<b>{f['label']}</b>", cell), Paragraph(str(f["value"]), cell)]
                   for f in report["filters"]]
    if filter_rows:
        t = Table(filter_rows, colWidths=[40 * mm, 120 * mm], hAlign="LEFT")
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, GRID), ("INNERGRID", (0, 0), (-1, -1), 0.25, GRID),
                               ("BACKGROUND", (0, 0), (0, -1), ZEBRA)]))
        story += [Paragraph("Selected filters", h2), t]

    if report["summary"]:
        items = report["summary"]
        per_row = 5
        grid = []
        for i in range(0, len(items), per_row):
            chunk = items[i:i + per_row]
            grid.append([Paragraph(f"<font size=7 color='#64748b'>{it['label']}</font><br/>"
                                   f"<font size=12><b>{_fmt(it['value'])}</b></font>", cell) for it in chunk]
                        + [""] * (per_row - len(chunk)))
        width = (page[0] - 28 * mm) / per_row
        t = Table(grid, colWidths=[width] * per_row, hAlign="LEFT")
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, GRID), ("INNERGRID", (0, 0), (-1, -1), 0.5, GRID),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story += [Paragraph("Summary", h2), t]

    if report.get("chart"):
        drawing = _chart(report["chart"], page[0] - 28 * mm)
        if drawing:
            story += [Spacer(1, 4 * mm), drawing]

    usable = page[0] - 28 * mm
    for section in report["sections"]:
        columns = section["columns"]
        rows = section["rows"]
        heading = Paragraph(f"{section['title']} ({len(rows)})", h2)
        if not rows:
            story.append(KeepTogether([heading, Paragraph("No records match the selected filters.", small)]))
            continue
        data = [[Paragraph(c["label"], head) for c in columns]]
        for row in rows:
            data.append([Paragraph(_fmt(row.get(c["key"])), cell) for c in columns])
        weights = [2.6 if c["key"] in ("title", "finding", "department", "auditor", "lead_auditor", "reviewer",
                                       "owner") else 1.2 for c in columns]
        unit = usable / sum(weights)
        t = Table(data, colWidths=[w * unit for w in weights], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
            ("GRID", (0, 0), (-1, -1), 0.25, GRID),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [heading, t]

    story += [Spacer(1, 6 * mm), Paragraph(
        "Compliance scores are internal indicators based on recorded checklist responses and are not proof of legal "
        "or certification compliance.", small)]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(14 * mm, 8 * mm, f"AuditSphere — {report['title']} — Confidential")
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, f"Page {doc_.page}")
        canvas.setStrokeColor(ACCENT)
        canvas.setLineWidth(2)
        canvas.line(14 * mm, page[1] - 10 * mm, page[0] - 14 * mm, page[1] - 10 * mm)
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- Excel

HEADER_FILL = PatternFill("solid", fgColor="0F1B2D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _safe_sheet_title(title, used):
    base = "".join(ch for ch in title if ch not in r"[]:*?/\\")[:28] or "Sheet"
    name, i = base, 2
    while name in used:
        name = f"{base[:25]} {i}"
        i += 1
    used.add(name)
    return name


def _excel_value(value):
    """Prevent spreadsheet formula injection from user-entered text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def render_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    used = {"Summary"}
    ws["A1"] = report["organization"]
    ws["A1"].font = Font(size=10, color="64748B")
    ws["A2"] = report["title"]
    ws["A2"].font = Font(size=16, bold=True, color="0F1B2D")
    ws["A3"] = f"Generated {_generated(report)} by {report.get('generated_by', '')}"
    row = 5
    ws.cell(row=row, column=1, value="Selected filters").font = Font(bold=True)
    for f in report["filters"]:
        row += 1
        ws.cell(row=row, column=1, value=f["label"])
        ws.cell(row=row, column=2, value=_excel_value(str(f["value"])))
    row += 2
    ws.cell(row=row, column=1, value="Summary").font = Font(bold=True)
    for item in report["summary"]:
        row += 1
        ws.cell(row=row, column=1, value=item["label"])
        ws.cell(row=row, column=2, value=_excel_value(item["value"]))
    chart = report.get("chart")
    if chart and chart["labels"]:
        row += 2
        ws.cell(row=row, column=1, value=chart["title"]).font = Font(bold=True)
        row += 1
        ws.cell(row=row, column=1, value="Category").font = Font(bold=True)
        for j, s in enumerate(chart["series"], start=2):
            ws.cell(row=row, column=j, value=s["name"]).font = Font(bold=True)
        for i, label in enumerate(chart["labels"]):
            row += 1
            ws.cell(row=row, column=1, value=label)
            for j, s in enumerate(chart["series"], start=2):
                ws.cell(row=row, column=j, value=s["values"][i])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40

    for section in report["sections"]:
        sheet = wb.create_sheet(_safe_sheet_title(section["title"], used))
        columns = section["columns"]
        for j, c in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=j, value=c["label"])
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        for i, r in enumerate(section["rows"], start=2):
            for j, c in enumerate(columns, start=1):
                sheet.cell(row=i, column=j, value=_excel_value(r.get(c["key"])))
        for j, c in enumerate(columns, start=1):
            longest = max([len(str(c["label"]))] + [len(str(r.get(c["key"]) or "")) for r in section["rows"][:200]])
            sheet.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 60)
        sheet.freeze_panes = "A2"
        if section["rows"]:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(section['rows']) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
