import io
import zipfile

from django.test import TestCase
from openpyxl import load_workbook

from apps.activity_logs.models import AuditActivityLog
from apps.audits import services as audit_services
from apps.core.tests.factories import World, client_for
from apps.reports.builders import REPORTS


class ReportTests(TestCase):
    def setUp(self):
        self.w = World(start=True)
        sales_audit = audit_services.create_audit(self.w.audit_admin, {
            "title": "Sales audit", "department": self.w.other_department,
            "template_version": self.w.template.current_approved_version, "planned_start_date": "2026-01-01",
            "due_date": "2026-01-20", "lead_auditor": self.w.other_auditor, "reviewer": self.w.reviewer})
        audit_services.schedule(self.w.audit_admin, sales_audit)

    def test_catalog_and_every_report_builds(self):
        client = client_for(self.w.audit_admin)
        self.assertEqual(len(client.get("/api/reports/").json()), len(REPORTS))
        for key in REPORTS:
            body = client.get(f"/api/reports/{key}/").json()
            self.assertEqual(body["key"], key)
            self.assertEqual(body["organization"], self.w.org.name)
            self.assertTrue(body["sections"])

    def test_pdf_and_excel_exports(self):
        client = client_for(self.w.audit_admin)
        pdf = client.get("/api/reports/audits/export/pdf/?department=%d" % self.w.department.id)
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertIn("attachment", pdf["Content-Disposition"])
        xlsx = client.get("/api/reports/audits/export/excel/")
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(xlsx.content)))
        workbook = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames[0], "Summary")
        audits_sheet = workbook["Audits"]
        self.assertEqual(audits_sheet.max_row, 3)  # header + 2 audits
        log = AuditActivityLog.objects.filter(action="report.exported").first()
        self.assertEqual(log.changes["format"], "excel")

    def test_filters_are_applied_and_described(self):
        body = client_for(self.w.audit_admin).get(f"/api/reports/audits/?department={self.w.department.id}").json()
        self.assertEqual(len(body["sections"][0]["rows"]), 1)
        self.assertIn({"label": "Department", "value": "Finance"}, body["filters"])

    def test_report_data_is_scoped_to_the_user(self):
        body = client_for(self.w.manager).get("/api/reports/audits/").json()
        self.assertEqual({r["department"] for r in body["sections"][0]["rows"]}, {"Finance"})

    def test_action_owner_cannot_access_reports(self):
        client = client_for(self.w.owner)
        self.assertEqual(client.get("/api/reports/audits/").status_code, 403)
        self.assertEqual(client.get("/api/reports/audits/export/pdf/").status_code, 403)

    def test_invalid_filters_rejected(self):
        response = client_for(self.w.audit_admin).get("/api/reports/audits/?date_from=2026-05-01&date_to=2026-01-01")
        self.assertEqual(response.status_code, 400)

    def test_unknown_report_is_404(self):
        self.assertEqual(client_for(self.w.audit_admin).get("/api/reports/nope/").status_code, 404)


class DashboardTests(TestCase):
    def setUp(self):
        self.w = World(start=True)

    def test_dashboard_endpoints(self):
        client = client_for(self.w.executive)
        summary = client.get("/api/dashboard/summary/").json()
        self.assertEqual(summary["total_audits"], 1)
        self.assertEqual(summary["in_progress_audits"], 1)
        self.assertIn("scoring_explanation", summary)
        for path in ("status-distribution", "department-compliance", "risk-distribution", "monthly-trend",
                     "upcoming-audits", "recent-activity", "overdue-actions"):
            self.assertEqual(client.get(f"/api/dashboard/{path}/").status_code, 200, path)
        self.assertEqual(len(client.get("/api/dashboard/monthly-trend/").json()), 12)
